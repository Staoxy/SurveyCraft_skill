"""survey_export 的 golden / 冒烟测试。"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from common import ROOT, load_survey
from survey_export import export_wjx_text, export_word

DEMO = ROOT / "surveys" / "demo" / "survey.json"
GOLDEN = ROOT / "tests" / "golden" / "demo_wjx_text.txt"


def test_wjx_text_matches_golden():
    survey = load_survey(DEMO)
    text, warnings = export_wjx_text(survey)
    assert GOLDEN.exists(), "golden 文件缺失，请先生成"
    assert text == GOLDEN.read_text(encoding="utf-8")


def test_wjx_text_logic_warnings():
    survey = load_survey(DEMO)
    text, warnings = export_wjx_text(survey)
    # 6 条 visibleIf 逻辑 + 矩阵题量表标签 + 其他填空 等都会产生警告
    assert len(warnings) >= 6
    assert any("L1" in w and "逻辑设置" in w for w in warnings)
    assert any("Q09" in w and "矩阵" in w for w in warnings)


def test_wjx_text_structure():
    survey = load_survey(DEMO)
    text, _ = export_wjx_text(survey)
    lines = text.splitlines()
    assert lines[0] == "大学生AI学习工具使用情况调查"
    assert "1、您的性别是[单选题]" in lines
    assert any(l.startswith("8、") and "[量表题]" in l for l in lines)
    assert any(l.startswith("9、") and "[矩阵量表题]" in l for l in lines)
    assert any(l.startswith("12、") and "[填空题]" in l for l in lines)


def test_word_export_smoke():
    """注：不用 pytest 的 tmp_path 夹具——本机 pytest 临时目录有权限问题。"""
    import shutil
    import tempfile

    from docx import Document

    survey = load_survey(DEMO)
    tmp = tempfile.mkdtemp(prefix="sc_export_")
    try:
        out = Path(tmp) / "demo.docx"
        export_word(survey, out)
        assert out.exists() and out.stat().st_size > 5000
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]
        assert any("大学生AI学习工具使用情况调查" in t for t in texts)
        assert any("您的性别是" in t for t in texts)
        # 矩阵题与量表题应转成表格
        assert len(doc.tables) >= 2
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
